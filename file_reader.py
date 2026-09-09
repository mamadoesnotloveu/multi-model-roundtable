"""解析 Word / PPT / Excel 文件为纯文本，供讨论使用。

支持：.docx / .pptx / .xlsx（新版 Office 格式）。
不支持：.doc / .ppt / .xls（老二进制格式），请先另存为新格式。
"""
from __future__ import annotations

import io


def read_file_text(filename: str, data: bytes) -> str:
    """根据扩展名解析文件，返回纯文本。"""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        return _read_docx(data)
    if name.endswith(".pptx"):
        return _read_pptx(data)
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    raise ValueError("不支持的文件类型：{}".format(filename))


def _read_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        texts.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        texts.append(" | ".join(cells))
        if texts:
            parts.append("【第 {} 页】\n".format(i) + "\n".join(texts))
    return "\n\n".join(parts)


def _read_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append("【工作表：{}】".format(ws.title))
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)
