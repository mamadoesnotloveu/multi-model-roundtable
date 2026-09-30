"""生成讨论成果报告（Word / PDF），两种格式排版结构一致。

内容仅包含：
  1. 元信息：选手、起草人、独立评委、讨论轮次、总讨论时长、各方 token 消耗；
  2. 成果正文：最终方案（评委裁决），其中的 markdown 表格会渲染成真正的表格。

不包含：原始话题、附件内容、过程讨论内容。
"""
from __future__ import annotations

import io
import os

# Word 统一字体（拉丁 + 东亚都设为同一字体，避免中英文混排粗细不一）
_FONT = "PingFang SC"


def _format_duration(seconds):
    if not seconds:
        return "未知"
    try:
        seconds = int(seconds)
    except (TypeError, ValueError):
        return str(seconds)
    if seconds < 60:
        return "{} 秒".format(seconds)
    m, s = divmod(seconds, 60)
    if m < 60:
        return "{} 分 {} 秒".format(m, s)
    h, m = divmod(m, 60)
    return "{} 小时 {} 分 {} 秒".format(h, m, s)


def _get_athletes(data):
    a = data.get("athletes")
    if isinstance(a, list) and a:
        return [str(x) for x in a]
    for rnd in data.get("per_round", []):
        if rnd:
            return [u.get("speaker", "") for u in rnd]
    return []


def _rounds_text(data):
    run = data.get("rounds_run")
    cap = data.get("num_rounds")
    if run is not None and cap is not None and int(run) < int(cap):
        return "{} 轮（上限 {} 轮）".format(run, cap)
    if run is not None:
        return "{} 轮".format(run)
    return "未知"


def _meta_items(data):
    athletes = _get_athletes(data)
    return [
        ("运动员（选手）", "、".join(athletes) if athletes else "未知"),
        ("起草人", data.get("drafter") or "未知"),
        ("独立评委", data.get("judge") or "未知"),
        ("讨论轮次", _rounds_text(data)),
        ("总讨论时长", _format_duration(data.get("duration_seconds"))),
    ]


def _usage_rows(data):
    usage = data.get("usage") or {}
    rows = []
    for label in sorted(usage.keys()):
        u = usage[label]
        if not isinstance(u, dict):
            continue
        inp = u.get("input", 0)
        out = u.get("output", 0)
        rows.append([str(label), str(inp), str(out), str(inp + out)])
    return rows


def _md_table_to_rows(lines):
    """把连续的 markdown 表格行解析成 rows(list[list[str]])；不是合法表格则返回 None。"""
    rows = []
    for ln in lines:
        s = ln.strip()
        inner = s[1:-1] if (s.startswith("|") and s.endswith("|")) else s.strip("|")
        rows.append([c.strip() for c in inner.split("|")])
    # 去掉分隔行（如 |---|---|）
    if len(rows) >= 2:
        sep = rows[1]
        if all(c.replace("-", "").replace(":", "").strip() == "" for c in sep):
            rows = [rows[0]] + rows[2:]
    # 去掉 **bold** 标记与空行
    rows = [[c.replace("**", "") for c in row] for row in rows]
    rows = [r for r in rows if any(c.strip() for c in r)]
    return rows if rows else None


def _parse_blocks(final):
    """把最终方案解析成块列表：[("p", text, bold), ("table", rows)]"""
    lines = (final or "").split("\n")
    blocks = []
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].rstrip()
        if s.strip().startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            rows = _md_table_to_rows(tbl)
            if rows:
                blocks.append(("table", rows))
                continue
        if s.strip():
            bold = s.lstrip().startswith(("#", "【"))
            text = s.strip()
            if text.startswith("#"):
                text = text.lstrip("#").strip()
            blocks.append(("p", text, bold))
        i += 1
    return blocks


def _register_font():
    """返回 reportlab 可用的中文字体名。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    for path in ("/Library/Fonts/Arial Unicode.ttf",
                 "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"):
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("CJK", path))
                return "CJK"
            except Exception:
                pass
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def _set_run_font(run, size=11, bold=False):
    """给 Word run 统一设置拉丁 + 东亚字体，保证中英文排版一致。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = _FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
    run.font.size = Pt(size)
    run.font.bold = bold


def render_word(data) -> bytes:
    """生成 Word（.docx）成果报告，返回字节。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("多模型圆桌讨论 · 成果报告"), size=16, bold=True)

    # 元信息
    for k, v in _meta_items(data):
        para = doc.add_paragraph()
        _set_run_font(para.add_run(k + "："), size=11, bold=True)
        _set_run_font(para.add_run(v), size=11, bold=False)

    # token 消耗表
    usage_rows = _usage_rows(data)
    if usage_rows:
        para = doc.add_paragraph()
        _set_run_font(para.add_run("各方 token 消耗"), size=13, bold=True)
        header = ["选手", "输入 tokens", "输出 tokens", "合计"]
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            _set_run_font(cell.paragraphs[0].add_run(h), size=10, bold=True)
        for row in usage_rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                _set_run_font(cells[i].paragraphs[0].add_run(v), size=10, bold=False)

    # 成果正文
    para = doc.add_paragraph()
    _set_run_font(para.add_run("最终方案（评委裁决）"), size=13, bold=True)
    for block in _parse_blocks(data.get("final", "")):
        if block[0] == "p":
            text, bold = block[1], block[2]
            para = doc.add_paragraph()
            _set_run_font(para.add_run(text), size=11, bold=bold)
        else:
            rows = block[1]
            cols = max(len(r) for r in rows) if rows else 1
            t = doc.add_table(rows=len(rows), cols=cols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    val = row[ci] if ci < len(row) else ""
                    cell = t.rows[ri].cells[ci]
                    _set_run_font(cell.paragraphs[0].add_run(val), size=10, bold=(ri == 0))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(data) -> bytes:
    """生成 PDF 成果报告，返回字节。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_font()

    def esc(t):
        if t is None:
            return ""
        t = str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return t.replace("\n", "<br/>")

    title = ParagraphStyle("title", fontName=font, fontSize=17, leading=25, alignment=1, spaceAfter=8)
    h1 = ParagraphStyle("h1", fontName=font, fontSize=13, leading=19, spaceBefore=12, spaceAfter=6,
                        textColor=colors.HexColor("#1a3d7c"))
    meta = ParagraphStyle("meta", fontName=font, fontSize=10.5, leading=18, spaceAfter=2)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=16, spaceAfter=4)

    story = [Paragraph("多模型圆桌讨论 · 成果报告", title)]
    for k, v in _meta_items(data):
        story.append(Paragraph("<b>{}</b>：{}".format(esc(k), esc(v)), meta))
    story.append(Spacer(1, 6))

    def _table_style(t, ncols):
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))

    usage_rows = _usage_rows(data)
    if usage_rows:
        story.append(Paragraph("各方 token 消耗", h1))
        table_data = [["选手", "输入 tokens", "输出 tokens", "合计"]] + usage_rows
        t = Table(table_data, colWidths=[45 * mm, 42 * mm, 42 * mm, 41 * mm])
        _table_style(t, 4)
        story.append(t)

    story.append(Paragraph("最终方案（评委裁决）", h1))
    usable = A4[0] - 40 * mm  # 页面宽度减左右边距
    for block in _parse_blocks(data.get("final", "")):
        if block[0] == "p":
            text, bold = block[1], block[2]
            story.append(Paragraph("<b>{}</b>".format(esc(text)) if bold else esc(text), body))
        else:
            rows = block[1]
            cols = max(len(r) for r in rows) if rows else 1
            w = usable / cols
            t = Table(rows, colWidths=[w] * cols, repeatRows=1)
            _table_style(t, cols)
            story.append(t)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    doc.build(story)
    return buf.getvalue()
